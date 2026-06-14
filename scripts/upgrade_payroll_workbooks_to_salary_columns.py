"""Replace legacy payroll amount columns in demo workbooks with salary columns."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = [
    ROOT / 'outputs' / 'upload_templates' / 'upay-payroll-bulk-upload-template.xlsx',
    *sorted((ROOT / 'outputs' / 'forecast_payroll_history').glob('*.xlsx')),
]


def upgrade(path: Path) -> None:
    workbook = load_workbook(path)
    payroll = workbook['Payroll']
    headers = [cell.value for cell in payroll[1]]
    if 'gross_salary' not in headers:
        amount_column = headers.index('amount') + 1
        payroll.insert_cols(amount_column + 1)
        payroll.cell(1, amount_column).value = 'basic_salary'
        payroll.cell(1, amount_column + 1).value = 'gross_salary'
        for row in range(2, payroll.max_row + 1):
            legacy_gross = payroll.cell(row, amount_column).value
            if legacy_gross is None:
                continue
            payroll.cell(row, amount_column).value = round(float(legacy_gross) * 0.60, 2)
            payroll.cell(row, amount_column + 1).value = legacy_gross
            payroll.cell(row, amount_column).number_format = '#,##0.00'
            payroll.cell(row, amount_column + 1).number_format = '#,##0.00'
        for column in (amount_column, amount_column + 1):
            payroll.cell(1, column).font = copy(payroll.cell(1, 1).font)
            payroll.cell(1, column).fill = copy(payroll.cell(1, 1).fill)
            payroll.cell(1, column).alignment = copy(payroll.cell(1, 1).alignment)
            payroll.column_dimensions[payroll.cell(1, column).column_letter].width = 16
    headers = [cell.value for cell in payroll[1]]
    if 'wallet_type' in headers:
        payroll.delete_cols(headers.index('wallet_type') + 1)
    payroll.auto_filter.ref = None
    payroll.tables.clear()
    if 'Instructions' in workbook.sheetnames:
        instructions = workbook['Instructions']
        instructions.cell(2, 1).value = 'Required columns'
        instructions.cell(2, 2).value = 'phone_number, basic_salary, and gross_salary are required for every payroll row.'
        instructions.cell(4, 1).value = 'Salary format'
        instructions.cell(4, 2).value = 'Enter positive BDT values without commas or currency symbols. Gross salary cannot be lower than basic salary.'
        instructions.cell(5, 1).value = 'Account verification'
        instructions.cell(5, 2).value = 'The backend verifies each phone number against the active upay account and approved employee roster.'
        if instructions.max_row >= 6:
            instructions.cell(6, 1).value = None
            instructions.cell(6, 2).value = None
    workbook.save(path)


for workbook_path in WORKBOOKS:
    upgrade(workbook_path)
    print(workbook_path)
