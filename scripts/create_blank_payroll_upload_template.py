"""Create the blank payroll XLSX served by the corporate upload page."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'frontend' / 'public' / 'templates' / 'payroll-upload-template.xlsx'
HEADERS = ('employee_code', 'employee_name', 'phone_number', 'department', 'basic_salary', 'gross_salary')


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    payroll = workbook.active
    payroll.title = 'Payroll'

    header_fill = PatternFill('solid', fgColor='2D3142')
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    thin_border = Border(bottom=Side(style='thin', color='DCE3EC'))
    widths = (18, 28, 18, 22, 16, 16)

    for index, (header, width) in enumerate(zip(HEADERS, widths), start=1):
        cell = payroll.cell(1, index, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
        payroll.column_dimensions[cell.column_letter].width = width

    payroll.row_dimensions[1].height = 24
    payroll.freeze_panes = 'A2'

    instructions = workbook.create_sheet('Instructions')
    instructions.column_dimensions['A'].width = 34
    instructions.column_dimensions['B'].width = 90
    instructions['A1'] = 'Payroll upload instructions'
    instructions['B1'] = 'Use the Payroll sheet only. Keep its header row exactly as provided.'
    instructions['A2'] = 'Required columns'
    instructions['B2'] = 'phone_number, basic_salary, and gross_salary are required for every payroll row.'
    instructions['A3'] = 'Phone number format'
    instructions['B3'] = 'Use an 11-digit Upay number beginning with 01. Keep the column formatted as text so the leading zero remains.'
    instructions['A4'] = 'Amount format'
    instructions['B4'] = 'Enter positive BDT basic and gross salaries without commas or currency symbols. Gross salary cannot be lower than basic salary.'
    instructions['A5'] = 'Account verification'
    instructions['B5'] = 'The backend verifies the uploaded phone number against the active upay account and this company’s approved employee roster.'
    instructions['A6'] = 'Payroll month'
    instructions['B6'] = 'Select the payroll month in the Corporate Assist upload page before choosing this file. Do not add a date column to the Payroll sheet.'

    for row in instructions.iter_rows(min_row=1, max_row=6, min_col=1, max_col=2):
        for cell in row:
            cell.font = Font(name='Arial', size=10, bold=cell.column == 1, color='2D3142')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = thin_border
    instructions['A1'].fill = header_fill
    instructions['A1'].font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    instructions['B1'].fill = header_fill
    instructions['B1'].font = Font(name='Arial', size=10, color='FFFFFF')
    for row_number in range(1, 7):
        instructions.row_dimensions[row_number].height = 30

    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
