"""Create ten sequential payroll workbooks for forecast and anomaly testing."""

from copy import copy
from datetime import date
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "forecast_payroll_history" / "fmcg-payroll-2026-08.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "forecast_payroll_history"

PERIODS = [
    date(2026, 9, 1), date(2026, 10, 1), date(2026, 11, 1), date(2026, 12, 1),
    date(2027, 1, 1), date(2027, 2, 1), date(2027, 3, 1), date(2027, 4, 1),
    date(2027, 5, 1), date(2027, 6, 1),
]

# Each scenario is deliberate test data. Salary multipliers are large enough to
# cross the application's established-employee anomaly thresholds.
SCENARIOS = {
    "2026-09": [("EMP-105", "High salary variance", "Gross salary intentionally raised to BDT 95,000.")],
    "2026-11": [("EMP-110", "Unregistered mobile account", "Mobile number intentionally changed to 01799999999.")],
    "2026-12": [("EMP-115", "Low salary variance", "Gross salary intentionally reduced to BDT 5,000.")],
    "2027-01": [("EMP-120", "Inactive mobile account", "Mobile number intentionally changed to known inactive account 01755556677.")],
    "2027-03": [("EMP-121", "High salary variance", "Gross salary intentionally raised to BDT 125,000.")],
    "2027-05": [
        ("EMP-123", "High salary variance", "Gross salary intentionally raised to BDT 135,000."),
        ("EMP-126", "Unregistered mobile account", "Mobile number intentionally changed to 01899999999."),
    ],
}


def apply_test_scenario(payroll, period_key: str) -> None:
    rows = {payroll.cell(row, 1).value: row for row in range(2, 32)}
    if period_key == "2026-09":
        payroll.cell(rows["EMP-105"], 5).value = 57_000
        payroll.cell(rows["EMP-105"], 6).value = 95_000
    elif period_key == "2026-11":
        payroll.cell(rows["EMP-110"], 3).value = "01799999999"
    elif period_key == "2026-12":
        payroll.cell(rows["EMP-115"], 5).value = 3_000
        payroll.cell(rows["EMP-115"], 6).value = 5_000
    elif period_key == "2027-01":
        payroll.cell(rows["EMP-120"], 3).value = "01755556677"
    elif period_key == "2027-03":
        payroll.cell(rows["EMP-121"], 5).value = 75_000
        payroll.cell(rows["EMP-121"], 6).value = 125_000
    elif period_key == "2027-05":
        payroll.cell(rows["EMP-123"], 5).value = 81_000
        payroll.cell(rows["EMP-123"], 6).value = 135_000
        payroll.cell(rows["EMP-126"], 3).value = "01899999999"


def add_scenario_sheet(workbook, period: date) -> None:
    if "Test Scenarios" in workbook.sheetnames:
        del workbook["Test Scenarios"]
    sheet = workbook.create_sheet("Test Scenarios")
    sheet.append(["Test status", "Employee ID", "Expected detection", "Intentional test change"])
    scenarios = SCENARIOS.get(period.strftime("%Y-%m"), [])
    if scenarios:
        for employee_code, detection, change in scenarios:
            sheet.append(["Contains intentional flaw", employee_code, detection, change])
    else:
        sheet.append(["Clean control month", "—", "No intentional anomaly", "Normal monthly salary movement only."])

    header_fill = PatternFill("solid", fgColor="2D3142")
    warning_fill = PatternFill("solid", fgColor="FFF1E8")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if scenarios:
                cell.fill = warning_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 24, "B": 18, "C": 28, "D": 58}.items():
        sheet.column_dimensions[column].width = width


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Missing source payroll workbook: {SOURCE}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    for period_index, period in enumerate(PERIODS, start=1):
        destination = OUTPUT_DIR / f"fmcg-payroll-{period:%Y-%m}.xlsx"
        copy2(SOURCE, destination)
        workbook = load_workbook(destination)
        payroll = workbook["Payroll"]
        instructions = workbook["Instructions"]

        # Apply modest, deterministic growth to clean monthly payroll values.
        growth = 1 + (period_index * 0.008)
        for row in range(2, 32):
            august_basic = float(payroll.cell(row, 5).value)
            august_gross = float(payroll.cell(row, 6).value)
            employee_adjustment = ((row + period_index) % 5 - 2) * 100
            gross_salary = max(1, round((august_gross * growth) + employee_adjustment, -2))
            basic_salary = min(gross_salary, round(august_basic * growth, -2))
            payroll.cell(row, 5).value = basic_salary
            payroll.cell(row, 6).value = gross_salary
            payroll.cell(row, 5).number_format = '#,##0'
            payroll.cell(row, 6).number_format = '#,##0'

        apply_test_scenario(payroll, period.strftime("%Y-%m"))
        scenarios = SCENARIOS.get(period.strftime("%Y-%m"), [])
        instructions.cell(7, 1).value = "Reference payroll period"
        instructions.cell(7, 2).value = f"{period:%B %Y}. This month is identified by the filename."
        instructions.cell(7, 2).comment = Comment("Reference period created for payroll forecast testing.", "OpenAI Codex")
        instructions.cell(8, 1).value = "Anomaly test status"
        instructions.cell(8, 2).value = (
            f"Contains {len(scenarios)} intentional AI/account-detection test case(s). See the Test Scenarios sheet."
            if scenarios else "Clean control month with no intentional test flaws."
        )
        for row in (7, 8):
            for cell in instructions[row]:
                cell.font = copy(instructions.cell(6, cell.column).font)
                cell.alignment = copy(instructions.cell(6, cell.column).alignment)

        add_scenario_sheet(workbook, period)
        workbook.save(destination)
        print(destination)


if __name__ == "__main__":
    main()
